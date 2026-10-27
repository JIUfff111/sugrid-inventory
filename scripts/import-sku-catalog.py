"""Preview or import the ASUS/ROG model catalog without changing inventory.

Use the bundled Python runtime with openpyxl:

    python scripts/import-sku-catalog.py --database data/inventory.db
    python scripts/import-sku-catalog.py --database data/inventory.db --apply

An explicit database path is required. The default operation is read-only.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = (
    PROJECT_DIR
    / "docs/sku-catalog/01-deliverables/asus/2026-10-26/sku-catalog-asus-2026-10-26.xlsx"
)
EXPECTED_MODELS = 60
EXPECTED_CONFIGURATIONS = 94
OPTION_COLUMNS = {
    "cpuOptions": "CPU选项",
    "ramOptions": "内存选项",
    "gpuOptions": "显卡选项",
    "storageOptions": "存储选项",
    "screenOptions": "屏幕选项",
}


@dataclass(frozen=True)
class Model:
    source_key: str
    brand: str
    model: str
    year: int
    confidence: str
    options: dict[str, list[str]]

    @property
    def key(self) -> tuple[str, str, str, int]:
        return ("笔记本电脑", self.brand.casefold(), self.model.casefold(), self.year)

    @property
    def sku_id(self) -> str:
        digest = hashlib.sha256("\0".join(map(str, self.key)).encode("utf-8")).hexdigest()[:24]
        return f"sku_catalog_asus_{digest}"


def normalized_text(value: object) -> str:
    return " ".join(str(value or "").split())


def valid_url(value: object, where: str) -> str:
    url = normalized_text(value)
    parsed = urlsplit(url)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise ValueError(f"{where}: missing or invalid source URL")
    return url


def rows_by_header(workbook, sheet_name: str, required: tuple[str, ...]):
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Missing worksheet: {sheet_name}")
    rows = workbook[sheet_name].iter_rows(values_only=True)
    headers = [normalized_text(value) for value in next(rows)]
    if len(headers) != len(set(headers)):
        raise ValueError(f"{sheet_name}: duplicate column heading")
    missing = set(required) - set(headers)
    if missing:
        raise ValueError(f"{sheet_name}: missing columns {sorted(missing)}")
    for row_number, row in enumerate(rows, 2):
        if any(value is not None and str(value).strip() for value in row):
            yield row_number, dict(zip(headers, row))


def parse_options(value: object, where: str) -> list[str]:
    try:
        parsed = json.loads(str(value or "[]"))
    except (TypeError, ValueError) as error:
        raise ValueError(f"{where}: invalid JSON array") from error
    if not isinstance(parsed, list) or any(not isinstance(item, str) or not item.strip() for item in parsed):
        raise ValueError(f"{where}: expected an array of nonempty strings")
    options = [normalized_text(item) for item in parsed]
    if len({item.casefold() for item in options}) != len(options):
        raise ValueError(f"{where}: repeated option")
    return options


def source_model_key(brand: str, model: str, year: int) -> str:
    return f"笔记本电脑|{brand}|{model}|{year}"


def read_catalog(path: Path) -> tuple[list[Model], dict[str, int]]:
    if not path.is_file():
        raise ValueError(f"Workbook does not exist: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        models: list[Model] = []
        by_source_key: dict[str, Model] = {}
        model_columns = (
            "型号唯一键", "产品类型", "品牌", "型号", "型号年份", "显示名",
            "分类", *OPTION_COLUMNS.values(), "来源链接", "来源标题", "置信度",
        )
        for row_number, row in rows_by_header(workbook, "SKU型号", model_columns):
            where = f"SKU型号 row {row_number}"
            brand = normalized_text(row["品牌"])
            model = normalized_text(row["型号"])
            raw_year = row["型号年份"]
            if row["产品类型"] != "笔记本电脑" or brand not in {"华硕", "ROG"} or not model:
                raise ValueError(f"{where}: expected an ASUS/ROG laptop model")
            if isinstance(raw_year, bool) or str(raw_year).removesuffix(".0") not in {str(year) for year in range(2016, 2027)}:
                raise ValueError(f"{where}: invalid model year {raw_year!r}")
            year = int(raw_year)
            if model.casefold().startswith(brand.casefold()) or re.search(r"(?:^|\s)\d{4}(?:款)?$", model):
                raise ValueError(f"{where}: model must omit the separate brand and year")
            source_key = source_model_key(brand, model, year)
            if normalized_text(row["型号唯一键"]) != source_key:
                raise ValueError(f"{where}: model key does not match brand, model and year")
            if normalized_text(row["显示名"]) != f"{brand} {model} {year}":
                raise ValueError(f"{where}: display name repeats or omits brand/year")
            if row["分类"] not in {"轻薄本", "游戏本"}:
                raise ValueError(f"{where}: unexpected catalog category")
            confidence = normalized_text(row["置信度"]).lower()
            if confidence not in {"high", "medium"}:
                raise ValueError(f"{where}: low or missing confidence")
            valid_url(row["来源链接"], where)
            if not normalized_text(row["来源标题"]):
                raise ValueError(f"{where}: missing source title")
            options = {
                key: parse_options(row[column], f"{where} {column}")
                for key, column in OPTION_COLUMNS.items()
            }
            if any(not options[key] for key in ("cpuOptions", "gpuOptions", "storageOptions")):
                raise ValueError(f"{where}: missing required CPU, GPU or storage options")
            if source_key in by_source_key:
                raise ValueError(f"{where}: repeated model key {source_key}")
            item = Model(source_key, brand, model, year, confidence, options)
            models.append(item)
            by_source_key[source_key] = item

        if len(models) != EXPECTED_MODELS:
            raise ValueError(f"Expected {EXPECTED_MODELS} models, found {len(models)}")

        config_columns = ("组合编号", "型号唯一键", "显示名", "CPU", "内存", "显卡", "存储", "屏幕", "来源链接", "来源标题", "置信度")
        config_ids: set[str] = set()
        configs_per_model: Counter[str] = Counter()
        for row_number, row in rows_by_header(workbook, "配置组合", config_columns):
            where = f"配置组合 row {row_number}"
            config_id = normalized_text(row["组合编号"])
            model_key = normalized_text(row["型号唯一键"])
            if not config_id or config_id in config_ids or model_key not in by_source_key:
                raise ValueError(f"{where}: duplicate ID or unknown model key")
            if normalized_text(row["显示名"]) != f"{by_source_key[model_key].brand} {by_source_key[model_key].model} {by_source_key[model_key].year}":
                raise ValueError(f"{where}: display name does not match model")
            valid_url(row["来源链接"], where)
            if not normalized_text(row["来源标题"]) or normalized_text(row["置信度"]).lower() not in {"high", "medium"}:
                raise ValueError(f"{where}: missing source title or low confidence")
            config_ids.add(config_id)
            configs_per_model[model_key] += 1
        if len(config_ids) != EXPECTED_CONFIGURATIONS:
            raise ValueError(f"Expected {EXPECTED_CONFIGURATIONS} configurations, found {len(config_ids)}")
        models_without_configurations = sorted(set(by_source_key) - set(configs_per_model))

        alias_columns = ("别名", "型号唯一键")
        alias_count = 0
        for row_number, row in rows_by_header(workbook, "别名", alias_columns):
            if not normalized_text(row["别名"]) or normalized_text(row["型号唯一键"]) not in by_source_key:
                raise ValueError(f"别名 row {row_number}: empty alias or unknown model")
            alias_count += 1
        return models, {
            "configurations": len(config_ids),
            "aliases": alias_count,
            "modelsWithoutConfigurations": models_without_configurations,
            "missingRam": sum(not model.options["ramOptions"] for model in models),
            "missingScreen": sum(not model.options["screenOptions"] for model in models),
        }
    finally:
        workbook.close()


def canonical_model(brand: str, model: str, year: int | None) -> str:
    value = normalized_text(model)
    prefix = brand.casefold()
    if value.casefold().startswith(prefix):
        remainder = value[len(brand):]
        if brand != "ROG" or not remainder or not remainder[0].isalpha():
            value = remainder.lstrip(" -·") or value
    if year is not None:
        value = re.sub(rf"\s+{year}(?:款)?$", "", value).strip()
    return value.casefold()


def current_models(connection: sqlite3.Connection) -> dict[tuple[str, str, str, int], sqlite3.Row]:
    rows = connection.execute("""
        SELECT id, productType, brand, model, year,
               cpuOptions, ramOptions, gpuOptions, storageOptions, screenOptions
        FROM skus WHERE deleted = 0 AND productType = '笔记本电脑'
          AND brand IN ('华硕', 'ROG')
    """).fetchall()
    result = {}
    for row in rows:
        key = (row["productType"], row["brand"].casefold(), canonical_model(row["brand"], row["model"], row["year"]), row["year"])
        if key in result:
            raise ValueError(f"Database already has duplicate active ASUS/ROG SKUs: {key}")
        result[key] = row
    return result


def preview(connection: sqlite3.Connection, models: list[Model], details: dict[str, int]) -> dict:
    existing = current_models(connection)
    to_insert = [model for model in models if model.key not in existing]
    changed = []
    for model in models:
        previous = existing.get(model.key)
        if previous and any(json.loads(previous[column] or "[]") != values for column, values in model.options.items()):
            changed.append(model.source_key)
    return {
        "sourceModels": len(models),
        **details,
        "sourceByBrand": dict(Counter(model.brand for model in models)),
        "sourceByConfidence": dict(Counter(model.confidence for model in models)),
        "newModels": len(to_insert),
        "existingModels": len(models) - len(to_insert),
        "existingWithDifferentOptions": changed,
        "newModelKeys": [model.source_key for model in to_insert],
    }


def apply(connection: sqlite3.Connection, path: Path, models: list[Model], details: dict[str, int]) -> dict:
    plan = preview(connection, models, details)
    if not plan["newModels"]:
        return {"mode": "unchanged", **plan}
    backup = path.with_name(f"{path.stem}.before-asus-skus-{datetime.now():%Y%m%d-%H%M%S-%f}.db")
    destination = sqlite3.connect(backup)
    try:
        connection.backup(destination)
    finally:
        destination.close()
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    try:
        connection.execute("BEGIN IMMEDIATE")
        current = current_models(connection)
        actor = connection.execute("SELECT id FROM users WHERE role = 'admin' AND status = 'active' ORDER BY createdAt, id LIMIT 1").fetchone()
        if not actor:
            raise ValueError("No active administrator for SKU createdBy")
        inserted = 0
        for model in models:
            if model.key in current:
                continue
            connection.execute("""
                INSERT INTO skus (
                  id, productType, brand, model, year,
                  cpuOptions, ramOptions, gpuOptions, storageOptions, screenOptions,
                  priceMin, priceMax, deleted, createdBy, createdAt, updatedAt
                ) VALUES (?, '笔记本电脑', ?, ?, ?, ?, ?, ?, ?, ?, NULL, NULL, 0, ?, ?, ?)
            """, (
                model.sku_id, model.brand, model.model, model.year,
                *[json.dumps(model.options[key], ensure_ascii=False) for key in OPTION_COLUMNS],
                actor["id"], now, now,
            ))
            current[model.key] = None
            inserted += 1
        integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
        if integrity != "ok":
            raise ValueError(f"Post-import integrity_check failed: {integrity}")
        connection.commit()
    except Exception:
        connection.rollback()
        raise
    return {"mode": "applied", "inserted": inserted, "backup": str(backup), "integrity": integrity, **preview(connection, models, details)}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--database", type=Path, required=True, help="Existing SQLite database; required even for preview")
    parser.add_argument("--apply", action="store_true", help="Back up, then insert new models in a transaction")
    args = parser.parse_args()
    database = args.database.resolve()
    if not database.is_file():
        parser.error(f"Database does not exist: {database}")
    models, details = read_catalog(args.workbook)
    connection = sqlite3.connect(database if args.apply else f"{database.as_uri()}?mode=ro", timeout=30, uri=not args.apply)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA foreign_keys = ON")
    connection.execute("PRAGMA busy_timeout = 5000")
    try:
        before = preview(connection, models, details)
        print(json.dumps({"mode": "preview", "database": str(database), **before}, ensure_ascii=True, indent=2))
        if args.apply:
            print(json.dumps(apply(connection, database, models, details), ensure_ascii=True, indent=2))
    finally:
        connection.close()


if __name__ == "__main__":
    main()
