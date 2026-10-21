"""Import the Xiaomi/Redmi and ThinkPad model workbooks into the SKU library.

Run with the bundled Python runtime (which includes openpyxl):

    python scripts/import-laptop-skus.py --xiaomi <file.xlsx> --thinkpad <file.xlsx>
    python scripts/import-laptop-skus.py --xiaomi <file.xlsx> --thinkpad <file.xlsx> --apply

The first command previews the import. --apply makes a SQLite backup and writes
all changes in one transaction. Re-running --apply skips imported model names.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sqlite3
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT_DIR = Path(__file__).resolve().parents[1]
DEFAULT_DB = PROJECT_DIR / "data" / "inventory.db"
PRODUCT_TYPE = "笔记本电脑"
STORAGE_OPTIONS = ["256GB", "512GB", "1TB", "1.5TB", "2TB"]
MAIN_SHEET = "型号总表2016起"
OLDER_SHEET = "2016前及年份不明"


@dataclass(frozen=True)
class ModelRecord:
    source: str
    sheet: str
    row: int
    year: int | None
    brand: str
    model: str
    cpu: list[str]
    gpu: list[str]
    ram: list[str]


def normalize_year(value: object) -> int | None:
    """Keep a four digit release/model year when the workbook provides one."""
    if value is None:
        return None
    match = re.search(r"(?:19|20)\d{2}", str(value))
    return int(match.group(0)) if match else None


def choices(value: object) -> list[str]:
    if value is None:
        return []
    result: list[str] = []
    seen: set[str] = set()
    for raw in str(value).split(";"):
        item = " ".join(raw.split())
        if item and item.casefold() not in seen:
            result.append(item)
            seen.add(item.casefold())
    return result


def ram_capacity(token: str) -> int:
    """Return total installed GB from one MTM example, ignoring DDR speeds."""
    x = token.strip().replace("（", "(").replace("）", ")")
    reported = re.match(r"^(\d+)GB\s*\(([^)]*)\)", x, re.IGNORECASE)
    if reported:
        amounts = []
        for part in reported.group(2).split("+"):
            part = part.strip()
            suffix_count = re.fullmatch(r"(\d+)GB\s*[x×]\s*(\d+)", part, re.IGNORECASE)
            prefix_count = re.fullmatch(r"(\d+)\s*[x×]\s*(\d+)GB", part, re.IGNORECASE)
            single = re.fullmatch(r"(\d+)(?:GB)?", part, re.IGNORECASE)
            if suffix_count or prefix_count:
                match = suffix_count or prefix_count
                amounts.append(int(match.group(1)) * int(match.group(2)))
            elif single:
                amounts.append(int(single.group(1)))
            else:
                raise ValueError(f"Unrecognized RAM breakdown: {token!r}")
        if not amounts or sum(amounts) != int(reported.group(1)):
            raise ValueError(f"RAM total conflicts with components: {token!r}")
        return int(reported.group(1))

    total = 0
    for component in x.split("+"):
        component = component.strip()
        prefix_count = re.match(r"^(\d+)\s*[x×]\s*(\d+)GB\b", component, re.IGNORECASE)
        suffix_count = re.match(r"^(\d+)GB\s*[x×]\s*(\d+)\b", component, re.IGNORECASE)
        single = re.match(r"^(\d+)GB\b", component, re.IGNORECASE)
        if prefix_count:
            total += int(prefix_count.group(1)) * int(prefix_count.group(2))
        elif suffix_count:
            total += int(suffix_count.group(1)) * int(suffix_count.group(2))
        elif single:
            total += int(single.group(1))
        else:
            raise ValueError(f"Unrecognized RAM example: {token!r}")
    if total <= 0:
        raise ValueError(f"Invalid RAM total: {token!r}")
    return total


def ram_choices(value: object) -> list[str]:
    if value is None:
        return []
    sizes = {ram_capacity(x) for x in choices(value)}
    return [f"{size}GB" for size in sorted(sizes)]


def data_rows(workbook: Path, sheet_name: str):
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        sheet = book[sheet_name]
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if any(cell is not None for cell in values):
                yield row_number, values
    finally:
        book.close()


def read_xiaomi(workbook: Path) -> list[ModelRecord]:
    records = []
    for number, cells in data_rows(workbook, MAIN_SHEET):
        # Year, brand, family, model, status, CPU, GPU, RAM, ZOL count, note, URL.
        year, source_brand, _, model, _, cpu, gpu, ram, *_ = cells
        brand = {"小米": "小米", "红米": "Redmi"}.get(source_brand)
        if not brand:
            raise ValueError(f"Xiaomi row {number}: unexpected brand {source_brand!r}")
        records.append(ModelRecord("xiaomi", MAIN_SHEET, number, normalize_year(year), brand,
                                   str(model).strip(), choices(cpu), choices(gpu), ram_choices(ram)))
    return records


def read_thinkpad(workbook: Path) -> list[ModelRecord]:
    records = []
    for sheet_name in (MAIN_SHEET, OLDER_SHEET):
        for number, cells in data_rows(workbook, sheet_name):
            # Main: year, family, model, status, release, CPU, GPU, platform RAM,
            #       installed RAM examples, MTM count, year basis, PSREF key.
            # Older: year, family, model, status, CPU, GPU, year basis, PSREF key.
            if sheet_name == MAIN_SHEET:
                year, _, model, _, _, cpu, gpu, _, ram, *_ = cells
                memory_options = ram_choices(ram)
            else:
                year, _, model, _, cpu, gpu, *_ = cells
                memory_options = []  # The older sheet has no installed RAM examples.
            source_model = str(model).strip()
            if not source_model.startswith("ThinkPad "):
                raise ValueError(f"ThinkPad {sheet_name} row {number}: {source_model!r}")
            records.append(ModelRecord("thinkpad", sheet_name, number, normalize_year(year), "ThinkPad",
                                       source_model.removeprefix("ThinkPad "),
                                       choices(cpu), choices(gpu), memory_options))
    return records


def validate(records: list[ModelRecord]) -> None:
    if len(records) != 433:
        raise ValueError(f"Expected 433 source models, found {len(records)}")
    seen: set[tuple[str, str, int | None]] = set()
    for record in records:
        key = (record.brand, record.model, record.year)
        if key in seen:
            raise ValueError(f"Duplicate model in workbooks: {key}")
        seen.add(key)
        if not record.model or not record.cpu or not record.gpu:
            raise ValueError(f"Missing model/CPU/GPU at {record.source} {record.sheet} row {record.row}")
        if record.sheet == MAIN_SHEET and not record.ram:
            raise ValueError(f"Missing RAM at {record.source} row {record.row}")
    counts = Counter((r.source, r.sheet) for r in records)
    if counts != {("xiaomi", MAIN_SHEET): 65,
                  ("thinkpad", MAIN_SHEET): 330,
                  ("thinkpad", OLDER_SHEET): 38}:
        raise ValueError(f"Unexpected source sheet counts: {counts}")


def merged_options(previous: str | None, incoming: list[str]) -> str:
    current = json.loads(previous or "[]")
    existing = {str(value).casefold() for value in current}
    for option in incoming:
        if option.casefold() not in existing:
            current.append(option)
            existing.add(option.casefold())
    return json.dumps(current, ensure_ascii=False)


def stable_id(record: ModelRecord) -> str:
    key = f"{record.source}\0{record.sheet}\0{record.row}\0{record.brand}\0{record.model}\0{record.year}"
    return "sku_import_" + hashlib.sha256(key.encode("utf-8")).hexdigest()[:24]


def apply_records(connection: sqlite3.Connection, records: list[ModelRecord],
                  created_by: str, apply: bool) -> dict:
    current = connection.execute("""SELECT id, brand, model, year, cpuOptions, gpuOptions,
                                ramOptions, storageOptions FROM skus
                                WHERE deleted = 0 AND productType = ?""",
                                 (PRODUCT_TYPE,)).fetchall()
    by_name = {(row["brand"], row["model"], row["year"]): row for row in current}
    # Existing starter SKUs with a known one-to-one match in the source table.
    known_aliases = {
        ("ThinkPad", "X1 Carbon Gen 9"): ("ThinkPad", "X1 Carbon Gen9"),
    }
    result = {"sourceRows": len(records), "inserted": 0, "merged": 0,
              "unchanged": 0, "manualRam": 0, "brandsInserted": Counter(),
              "mergedModels": [], "nearMatchesKeptSeparate": [],
              "flaggedSourceModels": []}
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for record in records:
        if not record.ram:
            result["manualRam"] += 1
        if ".." in record.model:
            result["flaggedSourceModels"].append({"model": record.model,
                                                   "sourceRow": record.row})
        key = (record.brand, record.model, record.year)
        previous = by_name.get(key)
        if previous is None:
            # Existing imports from the first version did not persist year.
            # Use a NULL-year row as a one-time backfill candidate, but never
            # merge two explicitly different model years.
            previous = next((row for row in current
                             if row["brand"] == record.brand
                             and row["model"] == record.model
                             and row["year"] is None), None)
        if previous is None:
            alias = known_aliases.get((record.brand, record.model))
            if alias:
                previous = next((row for row in current
                                 if row["brand"] == alias[0]
                                 and row["model"] == alias[1]
                                 and (row["year"] is None or row["year"] == record.year)), None)
        if previous is not None:
            # Merge only exact names and the explicit ThinkPad alias. The
            # Redmi/Xiaomi starter models are similar names with conflicting
            # generations, so their workbook rows remain distinct SKUs.
            cpu = merged_options(previous["cpuOptions"], record.cpu)
            gpu = merged_options(previous["gpuOptions"], record.gpu)
            ram = merged_options(previous["ramOptions"], record.ram)
            storage = json.dumps(STORAGE_OPTIONS, ensure_ascii=False)
            options_unchanged = (cpu, gpu, ram, storage) == (previous["cpuOptions"], previous["gpuOptions"],
                                                             previous["ramOptions"], previous["storageOptions"])
            year_needs_backfill = previous["year"] is None and record.year is not None
            if options_unchanged and not year_needs_backfill:
                result["unchanged"] += 1
            else:
                if options_unchanged:
                    result["merged"] += 1
                else:
                    result["merged"] += 1
                    result["mergedModels"].append({"source": record.model,
                                                   "existing": previous["model"]})
                if apply:
                    connection.execute("""UPDATE skus SET year=?, cpuOptions=?, gpuOptions=?,
                        ramOptions=?, storageOptions=?, updatedAt=? WHERE id=?""",
                                       (record.year, cpu, gpu, ram, storage, now, previous["id"]))
            continue

        if key[:2] in {("Redmi", "Redmi Book 14"), ("Redmi", "Redmi Book Pro 15"),
                       ("小米", "Xiaomi Book Air 13")}:
            result["nearMatchesKeptSeparate"].append(record.model)
        result["inserted"] += 1
        result["brandsInserted"][record.brand] += 1
        if apply:
            connection.execute("""INSERT INTO skus (
                id, productType, brand, model, year, cpuOptions, ramOptions, gpuOptions,
                storageOptions, screenOptions, priceMin, priceMax, deleted,
                createdBy, createdAt, updatedAt
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, NULL, NULL, 0, ?, ?, ?)""",
                               (stable_id(record), PRODUCT_TYPE, record.brand, record.model, record.year,
                                json.dumps(record.cpu, ensure_ascii=False),
                                json.dumps(record.ram, ensure_ascii=False),
                                json.dumps(record.gpu, ensure_ascii=False),
                                json.dumps(STORAGE_OPTIONS, ensure_ascii=False), "[]",
                                created_by, now, now))
        # Keep the preview and repeated import lookup consistent within this run.
        row = {"id": stable_id(record), "brand": record.brand, "year": record.year,
                        "model": record.model,
                        "cpuOptions": json.dumps(record.cpu, ensure_ascii=False),
                        "gpuOptions": json.dumps(record.gpu, ensure_ascii=False),
                        "ramOptions": json.dumps(record.ram, ensure_ascii=False),
                        "storageOptions": json.dumps(STORAGE_OPTIONS, ensure_ascii=False)}
        by_name[key] = row
        current.append(row)
    result["brandsInserted"] = dict(result["brandsInserted"])
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xiaomi", type=Path, required=True)
    parser.add_argument("--thinkpad", type=Path, required=True)
    parser.add_argument("--database", type=Path, default=DEFAULT_DB)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    records = read_xiaomi(args.xiaomi) + read_thinkpad(args.thinkpad)
    validate(records)
    connection = sqlite3.connect(args.database, timeout=30)
    connection.row_factory = sqlite3.Row
    try:
        user = connection.execute("SELECT id FROM users ORDER BY createdAt LIMIT 1").fetchone()
        if not user:
            raise ValueError("Database has no user for createdBy")
        preview = apply_records(connection, records, user["id"], apply=False)
        print(json.dumps({"mode": "preview", **preview}, ensure_ascii=False, indent=2))
        if not args.apply or preview["inserted"] + preview["merged"] == 0:
            return
        backup = args.database.with_name(
            f"{args.database.stem}.before-laptop-skus-{datetime.now():%Y%m%d-%H%M%S}.db")
        destination = sqlite3.connect(backup)
        try:
            connection.backup(destination)
        finally:
            destination.close()
        print(f"Backup: {backup}")
        try:
            connection.execute("BEGIN IMMEDIATE")
            outcome = apply_records(connection, records, user["id"], apply=True)
            connection.commit()
        except Exception:
            connection.rollback()
            raise
        active = connection.execute("SELECT COUNT(*) FROM skus WHERE deleted=0").fetchone()[0]
        integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
        print(json.dumps({"mode": "applied", **outcome, "activeSkus": active,
                          "integrity": integrity}, ensure_ascii=False, indent=2))
    finally:
        connection.close()


if __name__ == "__main__":
    main()
